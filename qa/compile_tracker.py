#!/usr/bin/env python3
"""Compile qa/stories/*.json into the canonical feature tracker.

Canonical data: qa/user-stories.csv (one row per user story).
Rendered view:  qa/Feature-Tracker.xlsx (same rows, formatted, plus summary).

Each stories/*.json file is a JSON array of objects with keys:
  id, area, page, feature, story, expected, evidence, test_via
Tracker adds status columns which persist across recompiles (merged by id
from any existing user-stories.csv):
  status        BACKLOG | TESTED-PASS | TESTED-FAIL | FIXED | RETEST-PASS | RETEST-FAIL | BLOCKED | N/A
  error_found   description of any error observed during testing
  fix_ref       commit/short note for the fix
  notes
"""
import csv
import json
import sys
from pathlib import Path

QA = Path(__file__).parent
STORIES = QA / "stories"
CSV_PATH = QA / "user-stories.csv"
XLSX_PATH = QA / "Feature-Tracker.xlsx"

FIELDS = [
    "id", "area", "page", "feature", "story", "expected", "evidence",
    "test_via", "status", "error_found", "fix_ref", "notes",
]


def load_existing() -> dict:
    if not CSV_PATH.exists():
        return {}
    with CSV_PATH.open() as fh:
        return {row["id"]: row for row in csv.DictReader(fh)}


def main() -> None:
    existing = load_existing()
    rows = []
    seen = set()
    for path in sorted(STORIES.glob("*.json")):
        data = json.loads(path.read_text())
        for item in data:
            sid = item["id"].strip()
            if sid in seen:
                print(f"WARN duplicate id {sid} in {path.name}", file=sys.stderr)
                continue
            seen.add(sid)
            old = existing.get(sid, {})
            rows.append({
                "id": sid,
                "area": item.get("area", ""),
                "page": item.get("page", ""),
                "feature": item.get("feature", ""),
                "story": item.get("story", ""),
                "expected": item.get("expected", ""),
                "evidence": item.get("evidence", ""),
                "test_via": item.get("test_via", ""),
                "status": old.get("status", "BACKLOG"),
                "error_found": old.get("error_found", ""),
                "fix_ref": old.get("fix_ref", ""),
                "notes": old.get("notes", ""),
            })
    # keep rows whose story file vanished but that carry test state
    for sid, old in existing.items():
        if sid not in seen and old.get("status", "BACKLOG") != "BACKLOG":
            rows.append(old)

    rows.sort(key=lambda r: r["id"])
    with CSV_PATH.open("w", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {CSV_PATH} ({len(rows)} stories)")
    write_xlsx(rows)


def write_xlsx(rows) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ARIAL = "Arial"
    INK = "1B2530"
    thin = Side(style="thin", color="C9D1CD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    status_style = {
        "BACKLOG": ("5A6774", None),
        "TESTED-PASS": ("1E7F4F", "E4F2EA"),
        "RETEST-PASS": ("1E7F4F", "E4F2EA"),
        "FIXED": ("1F5FA8", "E8EFF8"),
        "TESTED-FAIL": ("B3442E", "F8E8E3"),
        "RETEST-FAIL": ("B3442E", "F8E8E3"),
        "BLOCKED": ("A06A14", "F7EEDA"),
        "N/A": ("8B96A0", "F2F4F2"),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Stories"
    ws.sheet_view.showGridLines = False
    for c, h in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FA8")
        cell.border = border
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, 2):
        for c, field in enumerate(FIELDS, 1):
            cell = ws.cell(row=r, column=c, value=row.get(field, ""))
            cell.font = Font(name=ARIAL, size=9, color=INK)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=field in
                                       ("story", "expected", "evidence", "error_found", "notes"))
        st = row.get("status", "BACKLOG")
        color, fill = status_style.get(st, (INK, None))
        scell = ws.cell(row=r, column=FIELDS.index("status") + 1)
        scell.font = Font(name=ARIAL, size=9, bold=True, color=color)
        if fill:
            scell.fill = PatternFill("solid", fgColor=fill)
    widths = dict(zip("ABCDEFGHIJKL", (10, 16, 26, 24, 42, 52, 30, 8, 13, 40, 12, 28)))
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{len(rows) + 1}"

    summary = wb.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    counts = {}
    areas = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
        areas[row["area"]] = areas.get(row["area"], 0) + 1
    summary.cell(row=1, column=1, value="Feature Tracker Summary").font = Font(name=ARIAL, size=13, bold=True)
    summary.cell(row=2, column=1, value=f"Total stories: {len(rows)}").font = Font(name=ARIAL, size=10)
    r = 4
    summary.cell(row=r, column=1, value="By status").font = Font(name=ARIAL, size=10, bold=True)
    for k in sorted(counts):
        r += 1
        summary.cell(row=r, column=1, value=k).font = Font(name=ARIAL, size=10)
        summary.cell(row=r, column=2, value=counts[k]).font = Font(name=ARIAL, size=10)
    r += 2
    summary.cell(row=r, column=1, value="By area").font = Font(name=ARIAL, size=10, bold=True)
    for k in sorted(areas):
        r += 1
        summary.cell(row=r, column=1, value=k).font = Font(name=ARIAL, size=10)
        summary.cell(row=r, column=2, value=areas[k]).font = Font(name=ARIAL, size=10)
    summary.column_dimensions["A"].width = 24

    # Audits sheet: one row per correctness/fidelity check from
    # qa/results/audit-*.json ({check, status, detail}).
    audit_files = sorted((QA / "results").glob("audit-*.json"))
    if audit_files:
        aud = wb.create_sheet("Audits")
        aud.sheet_view.showGridLines = False
        for c, h in enumerate(["audit", "check", "status", "detail"], 1):
            cell = aud.cell(row=1, column=c, value=h)
            cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5FA8")
            cell.border = border
        r = 2
        for path in audit_files:
            try:
                checks = json.loads(path.read_text())
            except Exception:
                continue
            for chk in checks:
                vals = [path.stem.replace("audit-", ""), chk.get("check", ""),
                        chk.get("status", ""), chk.get("detail", "")]
                for c, v in enumerate(vals, 1):
                    cell = aud.cell(row=r, column=c, value=v)
                    cell.font = Font(name=ARIAL, size=9, color=INK)
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=(c == 4))
                st = chk.get("status", "")
                scell = aud.cell(row=r, column=3)
                if st == "PASS":
                    scell.font = Font(name=ARIAL, size=9, bold=True, color="1E7F4F")
                    scell.fill = PatternFill("solid", fgColor="E4F2EA")
                elif st == "FAIL":
                    scell.font = Font(name=ARIAL, size=9, bold=True, color="B3442E")
                    scell.fill = PatternFill("solid", fgColor="F8E8E3")
                r += 1
        for col, w in zip("ABCD", (14, 40, 8, 100)):
            aud.column_dimensions[col].width = w
        aud.freeze_panes = "A2"
        aud.auto_filter.ref = f"A1:D{r - 1}"

    wb.save(XLSX_PATH)
    print(f"wrote {XLSX_PATH}")


if __name__ == "__main__":
    main()
